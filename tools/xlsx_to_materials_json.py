"""엑셀 DB(Pilot_DB 시트)를 앱이 읽는 materials.json으로 변환합니다.

사용법 (프로젝트 폴더에서):
    pip install openpyxl
    python tools/xlsx_to_materials_json.py

엑셀에서 원료를 추가·수정·삭제한 뒤 이 스크립트를 실행하고,
바뀐 materials.json을 GitHub에 올리면 앱에 반영됩니다.
"""
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data" / "materials_v2.0p_24_db.xlsx"
OUT = ROOT / "materials.json"
SHEET = "Pilot_DB"

# 엑셀 열 이름 -> materials.json 키 (기존 v2 materials.json과 같은 구조)
TEXT = {
    "Korean Name": "name_ko",
    "English Name": "name_en",
    "Latin Name": "latin_name",
    "Note": "note",
    "Category": "category",
    "Intensity": "intensity",
    "Description": "description",
    "Room Spray Tip": "room_spray_tip",
}
LISTS = {
    "Odor Family": "odor_family",
    "Effects": "effects",
    "Mood Keywords": "mood_keywords",
    "Safety Flags": "safety_flags",
}
ORDER = ["name_ko", "name_en", "latin_name", "note", "category", "odor_family",
         "effects", "mood_keywords", "safety_flags", "intensity", "description",
         "room_spray_tip"]


def split(value):
    return [x.strip() for x in str(value or "").split(",") if x.strip()]


def main(xlsx=XLSX, out=OUT):
    ws = openpyxl.load_workbook(xlsx, data_only=True)[SHEET]
    header = [c.value for c in ws[1]]
    missing = [h for h in list(TEXT) + list(LISTS) if h not in header]
    if missing:
        sys.exit(f"엑셀에 필요한 열이 없습니다: {missing}")
    col = {h: header.index(h) for h in header if h}

    materials = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[col["Korean Name"]]:
            continue
        item = {key: str(row[col[h]] or "").strip() for h, key in TEXT.items()}
        item.update({key: split(row[col[h]]) for h, key in LISTS.items()})
        if item["note"] not in ("top", "middle", "base"):
            sys.exit(f"{item['name_ko']}: Note 값이 top/middle/base가 아닙니다 ({item['note']})")
        materials.append({k: item[k] for k in ORDER})

    names = [m["name_ko"] for m in materials]
    dup = {n for n in names if names.count(n) > 1}
    if dup:
        sys.exit(f"원료명이 중복됩니다: {dup}")

    out.write_text(json.dumps(materials, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    notes = {n: sum(m["note"] == n for m in materials) for n in ("top", "middle", "base")}
    print(f"{out.name} 저장: 원료 {len(materials)}종 {notes}")


if __name__ == "__main__":
    main()
